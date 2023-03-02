from decimal import Decimal
import openpyxl
import tkinter as tk
import tkinter.font as f
from calculater import func as calculater
import re

class calculate_molecular_weight(tk.Frame):
    def __init__(self, master):
        super().__init__(master)
        self.master = master
        self.master.title("分子量計算ツール")
        self.master.geometry("600x300")

        self.font_large = f.Font(size=18)
        self.font_middle = f.Font(size=14)
        self.font_small = f.Font(size=12)
        self.font_middle_bold = f.Font(size=16, weight="bold")

        # フレームとボタンの作成
        self.create_text_frame()
        self.create_molecular_formula_frame()
        self.create_calculation_result_frame()

        self.text_frame.pack(pady=5)
        self.molecular_formula_frame.pack()
        self.calculation_result_frame.pack(pady=10)
    
    def create_text_frame(self):
        self.text_frame = tk.Frame(self.master)

        tk.Label(self.text_frame, 
                 text = "分子式を入力してEnterを押すと計算結果が表示されます。",
                 font = self.font_middle
                 ).pack(pady=10)
        
        tk.Label(self.text_frame,
                 text = "注意書き１",
                 font = self.font_small
                 ).pack()

        tk.Label(self.text_frame,
                 text = "注意書き２",
                 font = self.font_small
                 ).pack()

        tk.Label(self.text_frame,
                 text = "注意書き３",
                 font = self.font_small
                 ).pack()
    

    def create_molecular_formula_frame(self):
        self.molecular_formula_frame = tk.Frame(self.master)
        
        # ラベル
        tk.Label(self.molecular_formula_frame,
                 text = "分子式：",
                 font = self.font_middle_bold
                 ).pack(side = tk.LEFT)
        # テキストボックスの作成と配置
        self.molecular_formula_textbox = tk.Entry(
            self.molecular_formula_frame,
            width = 50
            )
        self.molecular_formula_textbox.pack(
            side = tk.LEFT
            )
        self.molecular_formula_textbox.insert(tk.END, "C3H8O")

        # Enter入力でcalculater起動
        self.molecular_formula_textbox.bind("<Return>", self.show_result)
    

    def create_calculation_result_frame(self):
        self.calculation_result_frame = tk.Frame(self.master)

        # ラベル
        tk.Label(self.calculation_result_frame,
                 text = "分子量計算結果",
                 font = self.font_large
                 ).pack()
        
        # 計算結果を表示する部分
        self.result_textbox = tk.Text(self.calculation_result_frame)
        self.result_textbox.tag_configure("center", justify="center")
        self.result_textbox.configure(
            font = self.font_large,
            pady = 20,
            state="disabled")
        self.result_textbox.pack()


    # 計算機部分
    # ランレングス圧縮された分子式から分子量を計算
    # ex) C3H8O => 60.095
    def show_result(self, event):
        formula = self.molecular_formula_textbox.get()
        self.result_textbox.configure(state="normal")
        self.result_textbox.delete("1.0", "end")
        self.result_textbox.insert("1.0", calculater(formula), "center")
        self.result_textbox.configure(state="disabled")



root = tk.Tk()
app = calculate_molecular_weight(master=root)
app.mainloop()